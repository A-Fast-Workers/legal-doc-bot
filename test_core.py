import asyncio
import io
import os
import unittest

from legalbot import prompts
from legalbot.analyzer import Analyzer
from legalbot.chunking import select_relevant, split_text
from legalbot.config import Settings, load_settings
from legalbot.errors import AccessDenied, BotError, ConfigError, DocumentError, RateLimited
from legalbot.extract import extract_text, normalize
from legalbot.formatting import split_message, to_html
from legalbot.llm import FakeLLM
from legalbot.ratelimit import RateLimiter
from legalbot.service import LegalService
from legalbot.storage import SessionStore

CONTRACT = ("1. Предмет договора. Исполнитель разрабатывает бота для Заказчика.\n" * 3).strip()
LONG_TEXT = "\n".join(f"{i}. Пункт номер {i}: стороны обязуются выполнять условия договора надлежащим образом." for i in range(1, 700))


FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
    "C:/Windows/Fonts/arial.ttf",
    "/Library/Fonts/Arial.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
]


def run(coro):
    return asyncio.run(coro)


class ConfigTests(unittest.TestCase):
    def test_missing_token(self):
        with self.assertRaises(ConfigError):
            load_settings({})

    def test_ok_and_ids(self):
        s = load_settings({"TELEGRAM_BOT_TOKEN": "t", "GIGACHAT_CREDENTIALS": "c", "ALLOWED_USER_IDS": "1, 2;3"})
        self.assertEqual(s.allowed_user_ids, frozenset({1, 2, 3}))
        self.assertFalse(s.gigachat_verify_ssl)

    def test_bad_int(self):
        with self.assertRaises(ConfigError):
            load_settings({"TELEGRAM_BOT_TOKEN": "t", "GIGACHAT_CREDENTIALS": "c", "MAX_FILE_MB": "abc"})


class ExtractTests(unittest.TestCase):
    def test_txt_utf8_and_cp1251(self):
        text = "Договор оказания услуг между сторонами. " * 5
        self.assertIn("Договор", extract_text(text.encode("utf-8"), "a.txt", 10_000).text)
        self.assertIn("Договор", extract_text(text.encode("cp1251"), "a.txt", 10_000).text)

    def test_too_short_and_too_long(self):
        with self.assertRaises(DocumentError):
            extract_text(b"abc", "a.txt", 1000)
        with self.assertRaises(DocumentError):
            extract_text(("слово " * 500).encode(), "a.txt", 1000)

    def test_unsupported(self):
        with self.assertRaises(DocumentError):
            extract_text(b"x", "a.jpg", 1000)
        with self.assertRaises(DocumentError):
            extract_text(b"x", "a.doc", 1000)

    def test_docx_with_table(self):
        from docx import Document
        d = Document()
        d.add_paragraph("Договор поставки между ООО «А» и ООО «Б». " * 3)
        t = d.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text = "Цена"
        t.rows[0].cells[1].text = "100 рублей"
        buf = io.BytesIO()
        d.save(buf)
        out = extract_text(buf.getvalue(), "d.docx", 10_000).text
        self.assertIn("Цена | 100 рублей", out)

    def test_pdf_roundtrip_and_scan(self):
        try:
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.ttfonts import TTFont
            from reportlab.pdfgen import canvas
        except ImportError:
            self.skipTest("reportlab не установлен")
        font = next((f for f in FONT_CANDIDATES if os.path.exists(f)), None)
        if font is None:
            self.skipTest("не найден TTF-шрифт с кириллицей")
        pdfmetrics.registerFont(TTFont("DJ", font))
        buf = io.BytesIO()
        c = canvas.Canvas(buf)
        c.setFont("DJ", 12)
        for i in range(3):
            c.drawString(50, 800 - i * 20, "Договор аренды помещения, стоимость 50000 рублей в месяц")
        c.save()
        doc = extract_text(buf.getvalue(), "d.pdf", 10_000)
        self.assertEqual(doc.pages, 1)
        self.assertIn("аренды", doc.text)

        buf2 = io.BytesIO()
        c2 = canvas.Canvas(buf2)
        c2.rect(10, 10, 50, 50)
        c2.save()
        with self.assertRaises(DocumentError):
            extract_text(buf2.getvalue(), "scan.pdf", 10_000)

    def test_normalize(self):
        self.assertEqual(normalize("до-\nговор\n\n\n\nконец"), "договор\n\nконец")


class ChunkTests(unittest.TestCase):
    def test_split_respects_limit_and_keeps_content(self):
        parts = split_text(LONG_TEXT, 3000)
        self.assertTrue(all(len(p) <= 3000 for p in parts))
        self.assertEqual("".join(p.replace("\n", "") for p in parts), LONG_TEXT.replace("\n", ""))

    def test_split_long_line(self):
        parts = split_text("слово " * 2000, 1000)
        self.assertTrue(all(len(p) <= 1000 for p in parts))

    def test_select_relevant_finds_penalty(self):
        chunks = [f"Раздел {i}. Общие положения и порядок взаимодействия." for i in range(10)]
        chunks[6] = "Раздел 6. Неустойка за просрочку оплаты составляет один процент в день."
        chosen = select_relevant(chunks, "какая неустойка за просрочку?", k=2)
        self.assertIn(chunks[6], chosen)


class FormattingTests(unittest.TestCase):
    def test_escape_and_markup(self):
        out = to_html("## Риски\n- **Штраф** <1%> & более\n* пункт")
        self.assertIn("<b>Риски</b>", out)
        self.assertIn("• <b>Штраф</b> &lt;1%&gt; &amp; более", out)
        self.assertIn("• пункт", out)

    def test_split_keeps_tags_balanced(self):
        text = "\n".join(f"<b>строка {i}</b> " + "а" * 80 for i in range(200))
        parts = split_message(text, 1000)
        self.assertTrue(all(len(p) <= 1000 for p in parts))
        for p in parts:
            self.assertEqual(p.count("<b>"), p.count("</b>"))

    def test_split_unclosed_tag_across_parts(self):
        text = "<b>" + "\n".join("слово " * 30 for _ in range(60)) + "</b>"
        parts = split_message(text, 800)
        for p in parts:
            self.assertEqual(p.count("<b>"), p.count("</b>"))
            self.assertLessEqual(len(p), 800)


class StorageLimiterTests(unittest.TestCase):
    def test_ttl(self):
        now = [0.0]
        st = SessionStore(60, clock=lambda: now[0])
        st.put(1, "t", "n")
        now[0] = 30
        self.assertIsNotNone(st.get(1))
        now[0] = 80
        self.assertIsNotNone(st.get(1))  # продлён обращением
        now[0] = 200
        self.assertIsNone(st.get(1))

    def test_rate_limit(self):
        now = [0.0]
        rl = RateLimiter(2, 3600, clock=lambda: now[0])
        rl.check(1)
        rl.check(1)
        with self.assertRaises(RateLimited):
            rl.check(1)
        rl.check(2)
        now[0] = 4000
        rl.check(1)

    def test_rate_limit_disabled(self):
        rl = RateLimiter(0)
        for _ in range(100):
            rl.check(1)


class AnalyzerTests(unittest.TestCase):
    def setUp(self):
        self.s = Settings(single_pass_chars=2000, chunk_chars=1500, qa_chunk_chars=800)

    def test_single_pass(self):
        llm = FakeLLM()
        res = run(Analyzer(llm, self.s).analyze(CONTRACT))
        self.assertEqual(res.mode, "single")
        self.assertEqual(len(llm.calls), 1)
        self.assertIn(CONTRACT[:50], llm.calls[0][1])

    def test_map_reduce(self):
        llm = FakeLLM(lambda sys_, user: "- кратко")
        progress = []

        async def prog(i, n):
            progress.append((i, n))

        res = run(Analyzer(llm, self.s).analyze(LONG_TEXT, prog))
        self.assertEqual(res.mode, "map-reduce")
        n = res.chunks_used
        self.assertGreater(n, 5)
        self.assertEqual(progress[-1], (n, n))
        self.assertEqual(len(llm.calls), n + 1)
        self.assertIn("ВЫЖИМКИ", llm.calls[-1][1])

    def test_map_reduce_terminates_when_summaries_do_not_shrink(self):
        llm = FakeLLM(lambda s, u: "я" * 4000)   # каждая выжимка огромная
        res = run(Analyzer(llm, self.s).analyze(LONG_TEXT))
        self.assertEqual(res.mode, "map-reduce")

    def test_qa_picks_relevant_chunk(self):
        text = LONG_TEXT + "\nОсобое условие: неустойка за просрочку составляет 5 процентов."
        llm = FakeLLM()
        run(Analyzer(llm, self.s).ask(text, "Какая неустойка за просрочку?"))
        user = llm.calls[0][1]
        self.assertIn("5 процентов", user)
        self.assertLess(len(user), 6000)


class ServiceTests(unittest.TestCase):
    def make(self, **kw):
        return LegalService(FakeLLM(), Settings(**kw))

    def test_flow(self):
        svc = self.make()
        svc.load_text(1, CONTRACT * 5)
        res = run(svc.analyze(1))
        self.assertTrue(res.text)
        self.assertTrue(run(svc.ask(1, "Кто исполнитель?")))
        self.assertTrue(run(svc.quick(1, "money")))
        self.assertTrue(svc.reset(1))
        with self.assertRaises(BotError):
            run(svc.ask(1, "ещё вопрос"))

    def test_no_document(self):
        with self.assertRaises(BotError):
            run(self.make().analyze(5))

    def test_access_denied(self):
        svc = self.make(allowed_user_ids=frozenset({7}))
        with self.assertRaises(AccessDenied):
            svc.load_text(1, CONTRACT * 5)
        svc.load_text(7, CONTRACT * 5)

    def test_file_too_big(self):
        svc = self.make(max_file_mb=1)
        with self.assertRaises(BotError):
            svc.load_file(1, b"x" * (1024 * 1024 + 1), "a.txt")

    def test_rate_limit_applies(self):
        svc = self.make(requests_per_hour=2)
        svc.load_text(1, CONTRACT * 5)
        run(svc.ask(1, "a?"))
        run(svc.ask(1, "b?"))
        with self.assertRaises(RateLimited):
            run(svc.ask(1, "c?"))

    def test_sessions_isolated(self):
        svc = self.make()
        svc.load_text(1, "Договор А. " * 30)
        with self.assertRaises(BotError):
            run(svc.ask(2, "что в договоре?"))

    def test_unknown_quick_action(self):
        svc = self.make()
        svc.load_text(1, CONTRACT * 5)
        with self.assertRaises(BotError):
            run(svc.quick(1, "nope"))

    def test_prompts_are_russian_and_have_disclaimer(self):
        self.assertIn("русском", prompts.SYSTEM)
        self.assertIn("не юридическая консультация", prompts.DISCLAIMER)


if __name__ == "__main__":
    unittest.main()
